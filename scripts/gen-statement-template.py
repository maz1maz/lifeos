# Builds public/templates/statement-template.xlsx — the blank form users fill and load with «📥 ورود از اکسل»
# (src/today/src/statementImport.js finds the sheets and columns by these exact header texts).
# Run: python3 scripts/gen-statement-template.py
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = 'public/templates/statement-template.xlsx'
NAVY, GOLD, INPUT, MUTED = '0B1F3A', 'C8A24A', 'FFF9E8', '64748B'
thin = Side(style='thin', color='CBD5E1')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)
right = Alignment(horizontal='right', vertical='center', wrap_text=True, readingOrder=2)
head_font, head_fill = Font(name='Tahoma', bold=True, color='FFFFFF', size=10), PatternFill('solid', fgColor=NAVY)
in_fill = PatternFill('solid', fgColor=INPUT)
base = Font(name='Tahoma', size=10)
UNITS = 'مترمربع,مترطول,مترمکعب,عدد,کیلوگرم,تن,دستگاه,سرویس,کنترات'
KINDS = 'فروش / تأمین,نصب (مشمول بیمه),سایر'

wb = Workbook()

# ── items + contract info ──
ws = wb.active; ws.title = 'صورت وضعیت پروژه'; ws.sheet_view.rightToLeft = True
ws.merge_cells('B1:H1'); ws['B1'] = 'فهرست آیتم‌های قرارداد'; ws['B1'].font = Font(name='Tahoma', bold=True, size=15, color=NAVY); ws['B1'].alignment = center
ws.row_dimensions[1].height = 34
info = [('نام پروژه :', 'B3'), ('موضوع قرارداد :', 'B4'), ('کارفرما :', 'B5'), ('شماره قرارداد :', 'F3'), ('تاریخ قرارداد :', 'F4')]
for label, cell in info:
    c = ws[cell]; c.value = label; c.font = Font(name='Tahoma', bold=True, size=10); c.alignment = right
    v = ws.cell(row=c.row, column=c.column + 1); v.fill = in_fill; v.border = box; v.font = base; v.alignment = right
ws.merge_cells('C3:E3'); ws.merge_cells('C4:E4'); ws.merge_cells('C5:E5')
ws['H4'] = '(مثلاً ۱۴۰۵/۰۲/۰۶)'; ws['H4'].font = Font(name='Tahoma', size=8, color=MUTED); ws['H4'].alignment = right
ws['F5'] = 'نام پیمانکار'; ws['G5'] = 'شرکت'
for c in (ws['F5'], ws['G5']): c.font = Font(name='Tahoma', bold=True, size=9, color=MUTED); c.alignment = center
ws['G6'].fill = in_fill; ws['G6'].border = box; ws['G6'].font = base; ws['F6'] = '(نام شرکت پیمانکار)'; ws['F6'].font = Font(name='Tahoma', size=8, color=MUTED); ws['F6'].alignment = right

H = 8
headers = ['شماره آیتم قرارداد', 'شرح آیتم قرارداد', 'نوع', 'مقدار', 'واحد', 'فی', 'مبلغ']
for i, h in enumerate(headers):
    c = ws.cell(row=H, column=2 + i, value=h); c.font = head_font; c.fill = head_fill; c.alignment = center; c.border = box
ws.row_dimensions[H].height = 30
ROWS = 40
first, last = H + 1, H + ROWS
for r in range(first, last + 1):
    n = r - H
    ws.cell(row=r, column=2, value=n if n <= 30 else f'مازاد بر قرارداد ({n - 30})')
    for col in range(2, 9):
        c = ws.cell(row=r, column=col); c.border = box; c.font = base; c.alignment = center if col != 3 else right
        if col in (3, 4, 5, 6, 7): c.fill = in_fill
    ws.cell(row=r, column=8, value=f'=IF(OR(E{r}="",G{r}=""),"",E{r}*G{r})').number_format = '#,##0'
    ws.cell(row=r, column=5).number_format = '#,##0.00'; ws.cell(row=r, column=7).number_format = '#,##0'
    if n > 30: ws.cell(row=r, column=2).font = Font(name='Tahoma', size=9, bold=True, color='8A6414')
tot, disc = last + 1, last + 2
ws.cell(row=tot, column=2, value='جمع کل'); ws.cell(row=tot, column=7, value=f'=SUMPRODUCT(--(LEFT(B{first}:B{last},5)<>"مازاد"),H{first}:H{last})')
ws.cell(row=disc, column=2, value='تخفیف'); ws.cell(row=disc, column=7).fill = in_fill
ws.cell(row=disc, column=8, value='← مبلغ تخفیف کل قرارداد (اختیاری، ریال)').font = Font(name='Tahoma', size=8, color=MUTED); ws.cell(row=disc, column=8).alignment = right
for r in (tot, disc):
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    for col in range(2, 8):
        c = ws.cell(row=r, column=col); c.border = box; c.font = Font(name='Tahoma', bold=True, size=10); c.alignment = center
    ws.cell(row=r, column=7).number_format = '#,##0'
for col, w in zip('ABCDEFGH', (2, 20, 52, 20, 12, 13, 18, 22)): ws.column_dimensions[col].width = w
ws.column_dimensions['H'].width = 26
dv_unit = DataValidation(type='list', formula1=f'"{UNITS}"', allow_blank=True); dv_kind = DataValidation(type='list', formula1=f'"{KINDS}"', allow_blank=True)
ws.add_data_validation(dv_unit); ws.add_data_validation(dv_kind)
dv_unit.add(f'F{first}:F{last}'); dv_kind.add(f'D{first}:D{last}')
ws.freeze_panes = f'A{first}'

# ── measurement sheet (optional) ──
m = wb.create_sheet('ریزمتر'); m.sheet_view.rightToLeft = True
m.merge_cells('A1:H1'); m['A1'] = 'ریزمتره (اختیاری — مقدار انجام‌شدهٔ هر آیتم در هر صورت وضعیت)'; m['A1'].font = Font(name='Tahoma', bold=True, size=13, color=NAVY); m['A1'].alignment = center
m['A3'] = 'دوره کارکرد :'; m['A3'].font = Font(name='Tahoma', bold=True); m['A3'].alignment = right; m['B3'].fill = in_fill; m['B3'].border = box; m.merge_cells('B3:C3')
m['D3'] = '(مثلاً ۱۴۰۵/۰۴/۲۲-۱۴۰۵/۰۶/۰۶ — دورهٔ آخرین صورت وضعیت)'; m['D3'].font = Font(name='Tahoma', size=8, color=MUTED); m['D3'].alignment = right
mh = ['آیتم قرارداد', 'شماره صورت وضعیت', 'تاریخ انجام', 'شرح آیتم قرارداد', 'مقدار انجام شده', 'فی', 'مبلغ', 'توضیحات']
for i, h in enumerate(mh):
    c = m.cell(row=5, column=1 + i, value=h); c.font = head_font; c.fill = head_fill; c.alignment = center; c.border = box
items = "'صورت وضعیت پروژه'!$B$%d:$G$%d" % (first, last)
for r in range(6, 206):
    for col in range(1, 9):
        c = m.cell(row=r, column=col); c.border = box; c.font = base; c.alignment = center if col != 4 and col != 8 else right
        if col in (1, 2, 3, 5, 8): c.fill = in_fill
    m.cell(row=r, column=4, value=f'=IF(A{r}="","",IFERROR(VLOOKUP(A{r},{items},2,0),"؟ آیتم پیدا نشد"))')
    m.cell(row=r, column=6, value=f'=IF(A{r}="","",IFERROR(VLOOKUP(A{r},{items},6,0),""))').number_format = '#,##0'
    m.cell(row=r, column=7, value=f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r})').number_format = '#,##0'
for col, w in zip('ABCDEFGH', (20, 14, 14, 46, 14, 16, 20, 36)): m.column_dimensions[col].width = w
m.freeze_panes = 'A6'

# ── guide ──
g = wb.create_sheet('راهنما'); g.sheet_view.rightToLeft = True; g.column_dimensions['A'].width = 110
lines = [
    ('راهنمای پر کردن فایل', True),
    ('۱. در برگهٔ «صورت وضعیت پروژه» فقط خانه‌های کرم‌رنگ را پر کنید؛ ستون «مبلغ» خودکار حساب می‌شود.', False),
    ('۲. ردیف‌های ۱ تا ۳۰ آیتم‌های اصلی قرارداد و ردیف‌های «مازاد بر قرارداد (n)» کارهای اضافه‌اند. ردیف خالی وارد نمی‌شود.', False),
    ('۳. «نوع»: فروش / تأمین (بدون بیمه) — نصب (مشمول بیمهٔ ۱۶٫۶۷٪) — سایر. اگر خالی بماند از روی شرح حدس زده می‌شود.', False),
    ('۴. «فی» و «مبلغ» به ریال. «مقدار» می‌تواند اعشاری باشد (مثلاً 249.89).', False),
    ('۵. «تخفیف»: مبلغ کل تخفیف قرارداد به ریال (اختیاری). درصد تخفیف خودکار از آن حساب می‌شود.', False),
    ('۶. برگهٔ «ریزمتر» اختیاری است: شمارهٔ آیتم (همان ستون اول برگهٔ قبل، مثل 3 یا «مازاد بر قرارداد (1)»)، شمارهٔ صورت وضعیت، تاریخ شمسی و مقدار انجام‌شده.', False),
    ('   اگر شمارهٔ صورت وضعیت را «قطعی» بنویسید، آن ردیف‌ها صورت وضعیت بعدی و قطعی می‌شوند.', False),
    ('برگهٔ «نمونهٔ پرشده» فقط برای دیدن است و هیچ‌وقت وارد نمی‌شود.', False),
    ('۷. در LifeOS: پروژه › اطلاعات مالی › آیتم‌های قرارداد › «📥 ورود از اکسل» و همین فایل را انتخاب کنید.', False),
    ('نام برگه‌ها و عنوان ستون‌ها را تغییر ندهید؛ ستون‌ها از روی عنوانشان پیدا می‌شوند.', True),
]
for i, (t, b) in enumerate(lines, start=1):
    c = g.cell(row=i * 2 - 1, column=1, value=t); c.font = Font(name='Tahoma', size=12 if b and i == 1 else 10.5, bold=b, color=NAVY if b else '0F172A'); c.alignment = right
# ── filled example (generic names) — a copy of the items sheet; the importer reads the first sheet that has items ──
ex = wb.copy_worksheet(ws); ex.title = 'نمونهٔ پرشده'; ex.sheet_view.rightToLeft = True
ex['B1'] = 'نمونهٔ پرشده — فقط برای دیدن؛ فرم خودتان را در برگهٔ «صورت وضعیت پروژه» پر کنید'
ex['B1'].font = Font(name='Tahoma', bold=True, size=12, color='8A6414')
for cell, v in (('C3', 'پروژهٔ نمونه'), ('C4', 'تهیه و نصب نمای آلومینیومی ساختمان نمونه'), ('C5', 'شرکت کارفرمای نمونه'), ('G3', 'NM-1405-01'), ('G4', '1405/01/15'), ('G6', 'شرکت پیمانکار نمونه')): ex[cell] = v
sample = [('تهیه و ساخت پنجره لولایی و فیکس', 'فروش / تأمین', 250, 'مترمربع', 30000000), ('فروش لوور آلومینیومی', 'فروش / تأمین', 400, 'مترطول', 16000000),
          ('فروش زیرسازی آلومینیومی نمای سنگی', 'فروش / تأمین', 300, 'مترمربع', 22000000), ('نصب پنجره لولایی و فیکس', 'نصب (مشمول بیمه)', 250, 'مترمربع', 1000000),
          ('نصب و اجرای لوور آلومینیومی', 'نصب (مشمول بیمه)', 400, 'مترطول', 1600000), ('نصب زیرسازی نمای سنگی', 'نصب (مشمول بیمه)', 300, 'مترمربع', 2200000)]
extra = [('مازاد نصب زیرسازی نمای سنگی', 'نصب (مشمول بیمه)', 40, 'مترمربع', 2200000), ('جرثقیل', 'سایر', 1, 'عدد', 30000000)]
for i, row in enumerate(sample): [ex.cell(row=first + i, column=3 + k, value=v) for k, v in enumerate(row)]
for i, row in enumerate(extra): [ex.cell(row=first + 30 + i, column=3 + k, value=v) for k, v in enumerate(row)]
ex.cell(row=disc, column=7, value=1200000000)
ex.sheet_properties.tabColor = GOLD
wb.move_sheet(ex, offset=-(len(wb.sheetnames) - 2))  # right after the blank form
wb.active = 0
wb.save(OUT)
print('wrote', OUT)
